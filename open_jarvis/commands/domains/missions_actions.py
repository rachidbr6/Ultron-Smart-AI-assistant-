"""Reads the user's Missions.docx to report which tasks are still open.

Tasks are Word list items; a task is considered done when any run in its
paragraph carries a highlight color (the user marks finished items bright
green). Non-list paragraphs are treated as section headings.
"""

from __future__ import annotations

import ctypes
import io
import os
import re
from dataclasses import dataclass
from pathlib import Path

from open_jarvis.security.jarvis_admin import format_actionable_message

try:
    import docx
    from docx.opc.exceptions import PackageNotFoundError
except ImportError:  # pragma: no cover - optional dependency
    docx = None
    PackageNotFoundError = OSError


@dataclass(frozen=True)
class MissionTask:
    section: str
    text: str
    done: bool


def missions_file_path() -> Path:
    configured = os.getenv("JARVIS_MISSIONS_FILE_PATH")
    if configured:
        return Path(configured)
    return Path.home() / "OneDrive" / "Desktop" / "Missions.docx"


def _read_locked_file(path: Path) -> bytes:
    """Read a file's bytes even while another app (e.g. Word) has it open.

    Requests the same read/write/delete share mode Explorer and PowerShell
    use, instead of Python's default exclusive-ish open() sharing.
    """

    generic_read = 0x80000000
    share_all = 0x1 | 0x2 | 0x4  # FILE_SHARE_READ | FILE_SHARE_WRITE | FILE_SHARE_DELETE
    open_existing = 3
    file_attribute_normal = 0x80
    invalid_handle_value = ctypes.c_void_p(-1).value

    kernel32 = ctypes.windll.kernel32
    handle = kernel32.CreateFileW(
        str(path), generic_read, share_all, None, open_existing, file_attribute_normal, None
    )
    if handle in (invalid_handle_value, 0, None):
        raise OSError(f"Could not open {path} for reading.")

    try:
        size = kernel32.GetFileSize(handle, None)
        buffer = ctypes.create_string_buffer(size)
        bytes_read = ctypes.c_ulong(0)
        if not kernel32.ReadFile(handle, buffer, size, ctypes.byref(bytes_read), None):
            raise OSError(f"Could not read {path}.")
        return buffer.raw[: bytes_read.value]
    finally:
        kernel32.CloseHandle(handle)


def _open_document(path: Path):
    """Open the docx, reading around another app's (e.g. Word's) file lock if needed."""

    try:
        return docx.Document(str(path))
    except (OSError, PackageNotFoundError):
        return docx.Document(io.BytesIO(_read_locked_file(path)))


def _is_list_item(paragraph) -> bool:
    p_pr = paragraph._p.pPr
    return p_pr is not None and p_pr.numPr is not None


def _is_done(paragraph) -> bool:
    return any(run.font.highlight_color is not None for run in paragraph.runs)


_DUE_DATE_PATTERN = re.compile(r"\b(\d{2}/\d{2})\b")


def _clean_section_name(text: str) -> str:
    """Trim parenthetical notes but keep the due date for a name that reads well aloud."""

    name = text.split("(")[0].strip(" -–—")
    due_date = _DUE_DATE_PATTERN.search(text)
    return f"{name}, due {due_date.group(1)}" if due_date else name


def parse_missions(path: Path | None = None) -> list[MissionTask]:
    """Return every task found in the missions document, done or not."""

    target = path or missions_file_path()
    document = _open_document(target)

    tasks: list[MissionTask] = []
    section = "General"
    for paragraph in document.paragraphs:
        text = paragraph.text.strip()
        if not text:
            continue
        if _is_list_item(paragraph):
            tasks.append(MissionTask(section=section, text=text, done=_is_done(paragraph)))
        else:
            section = _clean_section_name(text)
    return tasks


def remaining_tasks(path: Path | None = None) -> list[MissionTask]:
    return [task for task in parse_missions(path) if not task.done]


def build_remaining_summary(path: Path | None = None) -> str:
    """Build a concise spoken summary of what's left across all sections."""

    remaining = remaining_tasks(path)
    if not remaining:
        return "Every task in your missions document is marked complete, sir. Well done."

    by_section: dict[str, int] = {}
    for task in remaining:
        by_section[task.section] = by_section.get(task.section, 0) + 1

    section_bits = "; ".join(f"{count} in {section}" for section, count in by_section.items())
    next_task = remaining[0].text
    return (
        f"You have {len(remaining)} tasks left, sir: {section_bits}. "
        f"Next up: {next_task}"
    )


def _missing_missions_action_message() -> str:
    return format_actionable_message(
        "I can't read your missions document, sir.",
        f"No file was found at {missions_file_path()}.",
        "Set JARVIS_MISSIONS_FILE_PATH in your .env to point at the correct Word document.",
    )


def handle_missions_action(action: str, params: dict, context: dict) -> bool | None:
    """Handle the remaining-tasks lookup."""

    if action != "list_remaining_tasks":
        return None

    speak = context["speak"]
    logger = context["logger"]
    translate_text = context.get("translate_text")

    if docx is None:
        speak(
            format_actionable_message(
                "The document reader isn't installed, sir.",
                "python-docx is required to read your missions file.",
                "Install it with pip install python-docx.",
            )
        )
        return True

    target = missions_file_path()
    if not target.exists():
        speak(_missing_missions_action_message())
        return True

    try:
        summary = build_remaining_summary(target)
    except (OSError, ValueError, KeyError) as exc:
        logger.warning("Failed to read missions document: %s", exc)
        speak(
            format_actionable_message(
                "I couldn't read your missions document, sir.",
                str(exc),
                "Make sure it's a valid Word document and try again.",
            )
        )
        return True

    if translate_text is not None:
        translated = translate_text(summary, target_language="English", logger=logger)
        if translated:
            summary = translated
        else:
            logger.warning("Missions summary translation unavailable, speaking the original text.")

    speak(summary)
    return True
